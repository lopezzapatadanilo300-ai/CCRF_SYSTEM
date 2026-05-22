from urllib import request

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Sum, Count
from django.http import HttpResponse
from django.utils import timezone
from .models import Sede, Usuario, Miembro, MovimientoFinanciero, ReporteInventario, InventarioItem
import datetime


# ─── AUTH ────────────────────────────────────────────────────────────────────

def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            if user.is_active:
                login(request, user)
                return redirect('dashboard')
            else:
                messages.error(request, 'Tu cuenta está desactivada. Contacta al pastor central.')
        else:
            messages.error(request, 'Usuario o contraseña incorrectos.')
    return render(request, 'login.html')


def logout_view(request):
    logout(request)
    return redirect('login')


# ─── DASHBOARD ───────────────────────────────────────────────────────────────

@login_required
def dashboard_view(request):
    user = request.user
    if user.rol == 'pastor_central':
        sedes = Sede.objects.all()
        total_sedes = sedes.filter(activa=True).count()
        total_miembros = Miembro.objects.filter(estado='activo').count()
        total_usuarios = Usuario.objects.filter(is_active=True).exclude(rol='pastor_central').count()
        inventarios_pendientes = ReporteInventario.objects.filter(estado='enviado').count()
        context = {
            'sedes': sedes,
            'total_sedes': total_sedes,
            'total_miembros': total_miembros,
            'total_usuarios': total_usuarios,
            'inventarios_pendientes': inventarios_pendientes,
        }
        return render(request, 'dashboard_central.html', context)
    else:
        sede = user.sede
        total_miembros = Miembro.objects.filter(sede=sede, estado='activo').count() if sede else 0
        total_ingresos = MovimientoFinanciero.objects.filter(
            sede=sede,
            fecha__month=timezone.now().month,
            fecha__year=timezone.now().year
        ).exclude(tipo='gasto').aggregate(total=Sum('monto'))['total'] or 0
        total_gastos = MovimientoFinanciero.objects.filter(
            sede=sede,
            fecha__month=timezone.now().month,
            fecha__year=timezone.now().year,
            tipo='gasto'
        ).aggregate(total=Sum('monto'))['total'] or 0
        ultimos_movimientos = MovimientoFinanciero.objects.filter(sede=sede).order_by('-fecha')[:5]
        context = {
            'total_miembros': total_miembros,
            'total_ingresos': total_ingresos,
            'total_gastos': total_gastos,
            'ultimos_movimientos': ultimos_movimientos,
        }
        return render(request, 'dashboard_sede.html', context)

# ─── SEDES ───────────────────────────────────────────────────────────────────
def sedes_list(request):
    if request.user.rol != 'pastor_central':
        return redirect('dashboard')
    sedes = Sede.objects.all().order_by('nombre')
    return render(request, 'sedes.html', {'sedes': sedes})

def sede_form(request, pk=None):
    if request.user.rol != 'pastor_central':
        return redirect('dashboard')
    sede = get_object_or_404(Sede, pk=pk) if pk else None
    if request.method == 'POST':
        nombre = request.POST.get('nombre')
        municipio = request.POST.get('municipio')
        direccion = request.POST.get('direccion', '')
        activa = request.POST.get('activa') == 'on'
        if sede:
            sede.nombre = nombre; sede.municipio = municipio
            sede.direccion = direccion; sede.activa = activa; sede.save()
            messages.success(request, 'Sede actualizada.')
        else:
            Sede.objects.create(nombre=nombre, municipio=municipio, direccion=direccion, activa=activa)
            messages.success(request, 'Sede creada.')
        return redirect('sedes_list')
    return render(request, 'sedes.html', {'sede': sede, 'modo': 'sede_form'})

def usuarios_list(request):
    if request.user.rol != 'pastor_central':
        return redirect('dashboard')
    usuarios = Usuario.objects.exclude(rol='pastor_central').order_by('sede__nombre', 'username')
    return render(request, 'sedes.html', {'usuarios': usuarios, 'modo': 'usuarios'})
def crear_usuario(request):

    if request.user.rol != 'pastor_central':
        return redirect('dashboard')

    sedes = Sede.objects.filter(activa=True)

    if request.method == 'POST':

        username = request.POST.get('username')
        password = request.POST.get('password')
        first_name = request.POST.get('first_name', '')
        last_name = request.POST.get('last_name', '')
        rol = request.POST.get('rol', 'pastor_sede')
        sede_id = request.POST.get('sede')

        sede = get_object_or_404(Sede, pk=sede_id)

        Usuario.objects.create_user(
            username=username,
            password=password,
            first_name=first_name,
            last_name=last_name,
            rol=rol,
            sede=sede
        )

        messages.success(request, 'Usuario creado.')

        return redirect('usuarios_list')

    return render(request, 'sedes.html', {'sedes': sedes, 'modo': 'usuario_form'})
def usuario_form(request):
    if request.user.rol != 'pastor_central':
        return redirect('dashboard')
    sedes = Sede.objects.filter(activa=True)
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name')
        rol = request.POST.get('rol')
        sede_id = request.POST.get('sede')
        sede = get_object_or_404(Sede, pk=sede_id)
        if Usuario.objects.filter(username=username).exists():
            messages.error(request, 'Ese usuario ya existe.')
        else:
            Usuario.objects.create_user(username=username, password=password,
                first_name=first_name, last_name=last_name, rol=rol, sede=sede)
            messages.success(request, f'Usuario {username} creado.')
            return redirect('usuarios_list')
    return render(request, 'sedes.html', {'sedes': sedes, 'modo': 'usuario_form'})
 
 #---MIEMBROS-------

def miembros_list(request):
    user = request.user
    if user.rol == 'pastor_central':
        sede_id = request.GET.get('sede')
        miembros = Miembro.objects.all()
        if sede_id:
            miembros = miembros.filter(sede_id=sede_id)
        sedes = Sede.objects.filter(activa=True)
        context = {'miembros': miembros, 'sedes': sedes, 'sede_sel': sede_id}
    else:
        estado = request.GET.get('estado', '')
        miembros = Miembro.objects.filter(sede=user.sede)
        if estado:
            miembros = miembros.filter(estado=estado)
        context = {'miembros': miembros, 'estado_sel': estado}
    return render(request, 'miembros.html', context)

def miembro_form(request, pk=None):
    miembro = get_object_or_404(Miembro, pk=pk) if pk else None
    if request.method == 'POST':
        nombre_completo = request.POST.get('nombre_completo')
        telefono = request.POST.get('telefono', '')
        direccion = request.POST.get('direccion', '')
        ministerio = request.POST.get('ministerio', '')
        estado = request.POST.get('estado', 'activo')
        if request.user.rol == 'pastor_central':
            sede = get_object_or_404(Sede, pk=request.POST.get('sede'))
        else:
            sede = request.user.sede
        if miembro:
            miembro.nombre_completo = nombre_completo; miembro.telefono = telefono
            miembro.direccion = direccion; miembro.ministerio = ministerio
            miembro.estado = estado; miembro.sede = sede; miembro.save()
            messages.success(request, 'Miembro actualizado.')
        else:
            Miembro.objects.create(nombre_completo=nombre_completo, telefono=telefono,
                direccion=direccion, ministerio=ministerio, estado=estado, sede=sede)
            messages.success(request, 'Miembro registrado.')
        return redirect('miembros_list')
    sedes = Sede.objects.filter(activa=True) if request.user.rol == 'pastor_central' else None
    return render(request, 'miembros.html', {'miembro': miembro, 'sedes': sedes, 'modo': 'form'})

def miembro_toggle(request, pk):
    miembro = get_object_or_404(Miembro, pk=pk)

    if miembro.estado == 'activo':
        miembro.estado = 'inactivo'
        messages.success(request, 'Miembro desactivado.')
    else:
        miembro.estado = 'activo'
        messages.success(request, 'Miembro activado.')

    miembro.save()

    return redirect('miembros_list')

#----FINANZAS-------

def finanzas_list(request):
    user = request.user
    if user.rol == 'pastor_central':
        movimientos = MovimientoFinanciero.objects.all().order_by('-fecha')
        sedes = Sede.objects.filter(activa=True)
        sede_id = request.GET.get('sede')
        if sede_id:
            movimientos = movimientos.filter(sede_id=sede_id)
    else:
        movimientos = MovimientoFinanciero.objects.filter(sede=user.sede).order_by('-fecha')
        sedes = None
        sede_id = None
    tipo = request.GET.get('tipo')
    if tipo:
        movimientos = movimientos.filter(tipo=tipo)
    total = movimientos.filter(tipo__in=['diezmo','ofrenda','primicia','accion_gracias']).aggregate(total=Sum('monto'))['total'] or 0
    context = {'movimientos': movimientos, 'sedes': sedes, 'sede_sel': sede_id, 'tipo_sel': tipo, 'total': total}
    return render(request, 'finanzas.html', context)

def finanza_form(request):
    user = request.user
    sede = None if user.rol == 'pastor_central' else user.sede
    miembros = Miembro.objects.filter(estado='activo') if user.rol == 'pastor_central' else Miembro.objects.filter(sede=sede, estado='activo')
    if request.method == 'POST':
        tipo = request.POST.get('tipo')
        monto = request.POST.get('monto')
        fecha = request.POST.get('fecha')
        descripcion = request.POST.get('descripcion', '')
        miembro_id = request.POST.get('miembro')
        if user.rol == 'pastor_central':
            sede = get_object_or_404(Sede, pk=request.POST.get('sede'))
            miembros = Miembro.objects.filter(sede=sede, estado='activo')
        miembro = None
        if tipo == 'diezmo':
            if not miembro_id:
                messages.error(request, 'El diezmo requiere un miembro.')
                sedes = Sede.objects.filter(activa=True) if user.rol == 'pastor_central' else None
                return render(request, 'finanzas.html', {'miembros': miembros, 'sedes': sedes, 'modo': 'form'})
            miembro = get_object_or_404(Miembro, pk=miembro_id)
        MovimientoFinanciero.objects.create(sede=sede, tipo=tipo, monto=monto,
            fecha=fecha, descripcion=descripcion, miembro=miembro, registrado_por=user)
        messages.success(request, 'Movimiento registrado.')
        return redirect('finanzas_list')
    sedes = Sede.objects.filter(activa=True) if user.rol == 'pastor_central' else None
    return render(request, 'finanzas.html', {'miembros': miembros, 'sedes': sedes, 'modo': 'form'})

def finanza_eliminar(request, pk):
    movimiento = get_object_or_404(MovimientoFinanciero, pk=pk)

    movimiento.delete()

    messages.success(request, 'Movimiento eliminado.')

    return redirect('finanzas_list')

def finanza_editar(request, pk):
    movimiento = MovimientoFinanciero.objects.get(pk=pk)

    if request.method == 'POST':
        movimiento.tipo = request.POST.get('tipo')
        movimiento.monto = request.POST.get('monto')
        movimiento.fecha = request.POST.get('fecha')
        movimiento.descripcion = request.POST.get('descripcion')

        movimiento.save()

    return redirect('finanzas_list')

#-----INVENTARIO-------
def inventario_list(request):

    user = request.user

    if user.rol == 'pastor_central':

        reportes = ReporteInventario.objects.all().order_by(
            '-anio',
            '-mes'
        )

    else:

        reportes = ReporteInventario.objects.filter(
            sede=user.sede
        ).order_by(
            '-anio',
            '-mes'
        )

    return render(request, 'inventario.html', {
        'reportes': reportes
    })

def inventario_form(request):

    user = request.user

    if request.method == 'POST':

        mes = request.POST.get('mes')
        anio = request.POST.get('anio')

        if user.rol == 'pastor_central':

            sede = get_object_or_404(
                Sede,
                pk=request.POST.get('sede')
            )

        else:

            sede = user.sede

        existe = ReporteInventario.objects.filter(
            sede=sede,
            mes=mes,
            anio=anio
        ).exists()

        if existe:

            messages.error(
                request,
                'Ya existe un reporte para ese mes.'
            )

        else:

            ReporteInventario.objects.create(
                sede=sede,
                mes=mes,
                anio=anio,
                creado_por=user
            )

            messages.success(
                request,
                'Reporte creado correctamente.'
            )

            return redirect('inventario_list')

    # ← LÍNEAS QUE FALTABAN
    
    import datetime
    sedes = Sede.objects.filter(activa=True)
    meses = range(1, 13)
    year = datetime.date.today().year
    return render(request, 'inventario.html', {
        'sedes': sedes,
        'meses': meses,
        'year': year,
        'modo': 'form'
    })

def inventario_detalle(request, pk):

    reporte = get_object_or_404(
        ReporteInventario,
        pk=pk
    )

    movimientos = MovimientoFinanciero.objects.filter(
        sede=reporte.sede,
        fecha__month=reporte.mes,
        fecha__year=reporte.anio
    )

    resumenes = movimientos.values('tipo').annotate(
        total=Sum('monto')
    )

    total_ingresos = movimientos.exclude(tipo='gasto').aggregate(
        total=Sum('monto')
    )['total'] or 0

    total_gastos = movimientos.filter(tipo='gasto').aggregate(
        total=Sum('monto')
    )['total'] or 0

    total_general = total_ingresos - total_gastos

    return render(request, 'inventario.html', {
        'reporte': reporte,
        'resumenes': resumenes,
        'total_ingresos': total_ingresos,
        'total_gastos': total_gastos,
        'total_general': total_general,
        'modo': 'detalle'
    })
def inventario_enviar(request, pk):

    reporte = get_object_or_404(
        ReporteInventario,
        pk=pk
    )

    reporte.estado = 'enviado'

    reporte.save()

    messages.success(
        request,
        'Reporte enviado correctamente.'
    )

    return redirect(
        'inventario_detalle',
        pk=pk
    )


def inventario_revisar(request, pk):

    if request.user.rol != 'pastor_central':

        return redirect('dashboard')

    reporte = get_object_or_404(
        ReporteInventario,
        pk=pk
    )

    reporte.estado = 'revisado'

    reporte.save()

    messages.success(
        request,
        'Reporte revisado correctamente.'
    )

    return redirect(
        'inventario_detalle',
        pk=pk
    )

def reportes_view(request):
    user = request.user
    anio = int(request.GET.get('anio', datetime.date.today().year))
    movimientos = MovimientoFinanciero.objects.filter(fecha__year=anio) if user.rol == 'pastor_central' else MovimientoFinanciero.objects.filter(sede=user.sede, fecha__year=anio)
    sede_id = request.GET.get('sede')
    if sede_id and user.rol == 'pastor_central':
        movimientos = movimientos.filter(sede_id=sede_id)
    sedes = Sede.objects.filter(activa=True) if user.rol == 'pastor_central' else None
    resumen = movimientos.values('tipo').annotate(total=Sum('monto')).order_by('tipo')
    return render(request, 'reportes.html', {'resumen': resumen, 'anio': anio, 'sedes': sedes, 'sede_sel': sede_id})

def ranking_view(request):
    user = request.user
    anio = int(request.GET.get('anio', datetime.date.today().year))
    diezmos = MovimientoFinanciero.objects.filter(tipo='diezmo', fecha__year=anio) if user.rol == 'pastor_central' else MovimientoFinanciero.objects.filter(tipo='diezmo', sede=user.sede, fecha__year=anio)
    sede_id = request.GET.get('sede')
    if sede_id and user.rol == 'pastor_central':
        diezmos = diezmos.filter(sede_id=sede_id)
    sedes = Sede.objects.filter(activa=True) if user.rol == 'pastor_central' else None
    ranking = diezmos.values('miembro__nombre_completo', 'miembro__sede__nombre').annotate(total=Sum('monto')).order_by('-total')
    return render(request, 'reportes.html', {'ranking': ranking, 'anio': anio, 'sedes': sedes, 'sede_sel': sede_id, 'modo': 'ranking'})


@login_required
def exportar_pdf(request):
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
    except ImportError:
        messages.error(request, 'Instala reportlab: pip install reportlab')
        return redirect('reportes')
    user = request.user
    anio = int(request.GET.get('anio', datetime.date.today().year))
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="reporte_ccrf_{anio}.pdf"'
    p = canvas.Canvas(response, pagesize=letter)
    p.setFont("Helvetica-Bold", 16)
    p.drawString(50, 750, f"CCRF System — Reporte Financiero {anio}")
    p.setFont("Helvetica", 12)
    if user.rol == 'pastor_central':
        movimientos = MovimientoFinanciero.objects.filter(fecha__year=anio)
    else:
        movimientos = MovimientoFinanciero.objects.filter(sede=user.sede, fecha__year=anio)
    resumen = movimientos.values('tipo').annotate(total=Sum('monto')).order_by('tipo')
    y = 700
    p.drawString(50, y, "Resumen por tipo:")
    y -= 20
    for item in resumen:
        p.drawString(70, y, f"{item['tipo'].capitalize()}: ${item['total']:,.2f}")
        y -= 20
    total_general = movimientos.aggregate(total=Sum('monto'))['total'] or 0
    y -= 10
    p.setFont("Helvetica-Bold", 12)
    p.drawString(50, y, f"Total general: ${total_general:,.2f}")
    p.showPage()
    p.save()
    return response


@login_required
def exportar_excel(request):
    try:
        import openpyxl
    except ImportError:
        messages.error(request, 'Instala openpyxl: pip install openpyxl')
        return redirect('reportes')
    from openpyxl import Workbook
    user = request.user
    anio = int(request.GET.get('anio', datetime.date.today().year))
    wb = Workbook()
    ws = wb.active
    ws.title = f"Reporte {anio}"
    ws.append(['Fecha', 'Tipo', 'Monto', 'Miembro', 'Sede', 'Registrado por'])
    if user.rol == 'pastor_central':
        movimientos = MovimientoFinanciero.objects.filter(fecha__year=anio).select_related('miembro', 'sede', 'registrado_por')
    else:
        movimientos = MovimientoFinanciero.objects.filter(sede=user.sede, fecha__year=anio).select_related('miembro', 'sede', 'registrado_por')
    for m in movimientos:
        ws.append([
            str(m.fecha), m.tipo, float(m.monto),
            m.miembro.nombre_completo if m.miembro else '—',
            m.sede.nombre, m.registrado_por.username if m.registrado_por else '—'
        ])
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="reporte_ccrf_{anio}.xlsx"'
    wb.save(response)
    return response

def usuario_toggle(request, pk):
    if request.user.rol != 'pastor_central':
        return redirect('dashboard')

    usuario = get_object_or_404(Usuario, pk=pk)

    usuario.is_active = not usuario.is_active
    usuario.save()

    if usuario.is_active:
        messages.success(request, 'Usuario activado.')
    else:
        messages.success(request, 'Usuario desactivado.')

    return redirect('usuarios_list')

def inventario_exportar_pdf(request, pk):
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
    except ImportError:
        messages.error(request, 'Instala reportlab: pip install reportlab')
        return redirect('inventario_list')

    reporte = get_object_or_404(ReporteInventario, pk=pk)
    movimientos = MovimientoFinanciero.objects.filter(
        sede=reporte.sede,
        fecha__month=reporte.mes,
        fecha__year=reporte.anio
    )
    resumenes = movimientos.values('tipo').annotate(total=Sum('monto'))
    total_ingresos = movimientos.exclude(tipo='gasto').aggregate(total=Sum('monto'))['total'] or 0
    total_gastos = movimientos.filter(tipo='gasto').aggregate(total=Sum('monto'))['total'] or 0
    total_general = total_ingresos - total_gastos

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="reporte_{reporte.sede.nombre}_{reporte.mes}_{reporte.anio}.pdf"'

    p = canvas.Canvas(response, pagesize=letter)

    # Encabezado
    p.setFont("Helvetica-Bold", 18)
    p.drawString(50, 750, "CCRF System — Reporte de Inventario")
    p.setFont("Helvetica", 12)
    p.drawString(50, 725, f"Sede: {reporte.sede.nombre}")
    p.drawString(50, 705, f"Período: {reporte.mes}/{reporte.anio}")
    p.drawString(50, 685, f"Estado: {reporte.get_estado_display()}")

    # Línea separadora
    p.line(50, 670, 550, 670)

    # Encabezados tabla
    y = 650
    p.setFont("Helvetica-Bold", 12)
    p.drawString(50, y, "TIPO")
    p.drawString(400, y, "TOTAL")
    y -= 10
    p.line(50, y, 550, y)
    y -= 20

    # Filas
    tipo_nombres = {
        'diezmo': 'Diezmos',
        'ofrenda': 'Ofrendas',
        'primicia': 'Primicias',
        'accion_gracias': 'Acción de Gracias',
        'gasto': 'Gastos',
    }

    p.setFont("Helvetica", 12)
    for r in resumenes:
        nombre = tipo_nombres.get(r['tipo'], r['tipo'])
        p.drawString(50, y, nombre)
        p.drawRightString(550, y, f"${r['total']:,.0f}")
        y -= 25

    # Totales
    p.line(50, y, 550, y)
    y -= 20

    p.setFont("Helvetica", 12)
    p.setFillColorRGB(0.09, 0.4, 0.2)
    p.drawString(50, y, "Total ingresos:")
    p.drawRightString(550, y, f"+ ${total_ingresos:,.0f}")
    y -= 22

    p.setFillColorRGB(0.6, 0.1, 0.1)
    p.drawString(50, y, "Total gastos:")
    p.drawRightString(550, y, f"- ${total_gastos:,.0f}")
    y -= 22

    p.line(50, y, 550, y)
    y -= 20

    p.setFont("Helvetica-Bold", 14)
    if total_general >= 0:
        p.setFillColorRGB(0.09, 0.4, 0.2)
    else:
        p.setFillColorRGB(0.6, 0.1, 0.1)
    p.drawString(50, y, "TOTAL GENERAL:")
    p.drawRightString(550, y, f"${total_general:,.0f}")

    p.showPage()
    p.save()
    return response


def inventario_exportar_excel(request, pk):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        messages.error(request, 'Instala openpyxl: pip install openpyxl')
        return redirect('inventario_list')

    reporte = get_object_or_404(ReporteInventario, pk=pk)
    movimientos = MovimientoFinanciero.objects.filter(
        sede=reporte.sede,
        fecha__month=reporte.mes,
        fecha__year=reporte.anio
    )
    resumenes = movimientos.values('tipo').annotate(total=Sum('monto'))
    total_ingresos = movimientos.exclude(tipo='gasto').aggregate(total=Sum('monto'))['total'] or 0
    total_gastos = movimientos.filter(tipo='gasto').aggregate(total=Sum('monto'))['total'] or 0
    total_general = total_ingresos - total_gastos

    tipo_nombres = {
        'diezmo': 'Diezmos',
        'ofrenda': 'Ofrendas',
        'primicia': 'Primicias',
        'accion_gracias': 'Acción de Gracias',
        'gasto': 'Gastos',
    }

    wb = Workbook()
    ws = wb.active
    ws.title = f"Reporte {reporte.mes}-{reporte.anio}"

    # Título
    ws.merge_cells('A1:B1')
    ws['A1'] = f"CCRF System — Reporte {reporte.sede.nombre} {reporte.mes}/{reporte.anio}"
    ws['A1'].font = Font(bold=True, size=14)

    ws['A2'] = "Estado:"
    ws['B2'] = reporte.get_estado_display()

    ws.append([])

    # Encabezados
    ws.append(['TIPO', 'TOTAL'])
    ws['A4'].font = Font(bold=True)
    ws['B4'].font = Font(bold=True)
    ws['A4'].fill = PatternFill("solid", fgColor="1a4a7a")
    ws['B4'].fill = PatternFill("solid", fgColor="1a4a7a")
    ws['A4'].font = Font(bold=True, color="FFFFFF")
    ws['B4'].font = Font(bold=True, color="FFFFFF")

    # Datos
    for r in resumenes:
        nombre = tipo_nombres.get(r['tipo'], r['tipo'])
        ws.append([nombre, float(r['total'])])

    ws.append([])

    # Totales
    row_ingresos = ws.max_row + 1
    ws.append(['Total ingresos', float(total_ingresos)])
    ws.cell(row=row_ingresos, column=1).font = Font(bold=True, color="166534")
    ws.cell(row=row_ingresos, column=2).font = Font(bold=True, color="166534")

    row_gastos = ws.max_row + 1
    ws.append(['Total gastos', float(total_gastos)])
    ws.cell(row=row_gastos, column=1).font = Font(bold=True, color="991b1b")
    ws.cell(row=row_gastos, column=2).font = Font(bold=True, color="991b1b")

    row_general = ws.max_row + 1
    ws.append(['TOTAL GENERAL', float(total_general)])
    ws.cell(row=row_general, column=1).font = Font(bold=True, size=13)
    ws.cell(row=row_general, column=2).font = Font(bold=True, size=13)

    # Ancho columnas
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 20

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="reporte_{reporte.sede.nombre}_{reporte.mes}_{reporte.anio}.xlsx"'
    wb.save(response)
    return response
